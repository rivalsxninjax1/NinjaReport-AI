"""DOCX report generation from a CompiledReport.

Image-stretch guarantee: add_picture() is ALWAYS called with `width` only,
never `height` — python-docx auto-scales the other dimension to preserve
native aspect ratio when only one is given. The width itself is capped and
never-upscaled via generators.image_fit.compute_display_size().
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from generators.image_fit import DEFAULT_DPI, compute_display_size
from generators.report_compiler import CompiledReport

MAX_IMAGE_WIDTH_IN = 6.0


def generate_docx(report: CompiledReport, output_path: Path) -> Path:
    doc = Document()
    _add_cover_page(doc, report)
    _add_footer_with_page_numbers(doc)

    figure_number = 0
    for section in report.sections:
        doc.add_heading(section.title, level=1)

        for item in section.items:
            if item.item_type == "text":
                if item.caption:
                    doc.add_heading(item.caption, level=3)
                for paragraph_text in (item.text or "").split("\n"):
                    if paragraph_text.strip():
                        doc.add_paragraph(paragraph_text)

            elif item.item_type == "evidence":
                figure_number += 1
                _add_evidence_item(doc, item, figure_number)

            elif item.item_type == "finding":
                _add_finding_item(doc, item)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path


def _add_cover_page(doc: Document, report: CompiledReport) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(report.project_name)
    run.bold = True
    run.font.size = Pt(28)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run(f"{report.template_type.upper()} Report")
    sub_run.font.size = Pt(16)

    generated = doc.add_paragraph()
    generated.alignment = WD_ALIGN_PARAGRAPH.CENTER
    generated.add_run(f"Generated: {report.generated_at}").italic = True

    doc.add_page_break()


def _add_evidence_item(doc: Document, item, figure_number: int) -> None:
    if item.image_path is None or not item.image_path.exists():
        doc.add_paragraph(f"[Evidence file not available: {item.caption}]")
        return

    width_px, height_px = item.image_size_px or (0, 0)
    display = compute_display_size(width_px, height_px, MAX_IMAGE_WIDTH_IN, dpi=DEFAULT_DPI)

    # Only `width` is ever passed — this is what guarantees the aspect
    # ratio is preserved (python-docx computes height automatically).
    doc.add_picture(str(item.image_path), width=Inches(display.width_in))

    caption_para = doc.add_paragraph()
    caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_run = caption_para.add_run(f"Figure {figure_number}: {item.caption}")
    caption_run.italic = True
    caption_run.font.size = Pt(10)


def _add_finding_item(doc: Document, item) -> None:
    finding = item.finding
    doc.add_heading(finding.title, level=2)

    severity_text = finding.severity if finding.severity else "Not yet approved by reviewer"
    meta = doc.add_paragraph()
    meta.add_run("Severity: ").bold = True
    meta.add_run(f"{severity_text}\n")
    meta.add_run("Status: ").bold = True
    meta.add_run(f"{finding.status}\n")
    meta.add_run("Affected Asset: ").bold = True
    meta.add_run(finding.affected_asset or "Not specified")

    if finding.description:
        doc.add_heading("Description", level=3)
        doc.add_paragraph(finding.description)

    if finding.technical_impact:
        doc.add_heading("Technical Impact", level=3)
        doc.add_paragraph(finding.technical_impact)

    if finding.business_impact:
        doc.add_heading("Business Impact", level=3)
        doc.add_paragraph(finding.business_impact)

    if finding.reproduction_steps:
        doc.add_heading("Reproduction Steps", level=3)
        for step in finding.reproduction_steps:
            doc.add_paragraph(step, style="List Number")

    if finding.remediation:
        doc.add_heading("Remediation", level=3)
        doc.add_paragraph(finding.remediation)

    cvss_text = finding.cvss_vector if finding.cvss_vector else "Not provided"
    doc.add_paragraph().add_run(f"CVSS Vector: {cvss_text}").italic = True

    if finding.references_list:
        doc.add_heading("References", level=3)
        for ref in finding.references_list:
            doc.add_paragraph(ref, style="List Bullet")

    if finding.evidence_ids:
        doc.add_paragraph(f"Referenced evidence: {', '.join(finding.evidence_ids)}")


def _add_footer_with_page_numbers(doc: Document) -> None:
    """Adds a 'Page X of Y' footer using low-level field-code XML, since
    python-docx has no high-level API for dynamic page numbers."""
    section = doc.sections[0]
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    paragraph.add_run("Page ")
    _add_field_code(paragraph, "PAGE")
    paragraph.add_run(" of ")
    _add_field_code(paragraph, "NUMPAGES")


def _add_field_code(paragraph, field_code: str) -> None:
    run = paragraph.add_run()

    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = f" {field_code} "

    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_separate)
    run._r.append(fld_char_end)
